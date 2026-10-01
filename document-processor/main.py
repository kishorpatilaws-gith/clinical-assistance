"""
Document processing service.

A standalone process the .NET API calls over HTTP - not a library it references.
Keeping it separate is the point: PDF/DOCX extraction and embedding are Python's
strength, and the same service serves both Part 4 (extraction) and Part 5 (RAG).

Run:  uvicorn main:app --port 8000
"""

import io
import logging
import os
from contextlib import asynccontextmanager

import pymupdf  # PyMuPDF. NB: `import fitz` is the deprecated spelling.
from docx import Document as DocxDocument
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field

from dotenv import load_dotenv

# Explicit path: find_dotenv() walks the call stack and fails under some runners.
load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"))

import chunking
import embedding
import vectorstore

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("document-processor")

# The HF hub logs every file it resolves at INFO, which drowns out our own lines.
logging.getLogger("httpx").setLevel(logging.WARNING)
logging.getLogger("huggingface_hub").setLevel(logging.WARNING)


@asynccontextmanager
async def lifespan(app: FastAPI):
    # Load the model at startup rather than on the first request. Otherwise the
    # first clinician to upload a document pays ~3s of model loading, and /health
    # would report ready while the service could not actually embed anything.
    log.info("Warming embedding model ...")
    embedding.get_model()
    try:
        vectorstore.ensure_collection()
        log.info("Vector store ready: %s", vectorstore.stats())
    except Exception as exc:
        # Extraction still works without Qdrant; only indexing and search need it.
        log.warning("Vector store unavailable (%s). /index and /search will fail.", exc)
    log.info("Ready.")
    yield


app = FastAPI(
    title="Clinical document processor",
    description="Extracts, chunks and embeds synthetic clinical documents.",
    version="2.0.0",
    lifespan=lifespan,
)

MAX_BYTES = 30 * 1024 * 1024

PDF_TYPES = {"application/pdf"}
DOCX_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
}
TEXT_TYPES = {"text/plain", "text/markdown", "text/csv"}


# ----------------------------- models -----------------------------

class ExtractionResult(BaseModel):
    text: str
    page_count: int
    char_count: int
    detected_type: str


class HealthResult(BaseModel):
    status: str
    pymupdf_version: str
    embedding: dict
    vector_store: dict


class EmbedRequest(BaseModel):
    texts: list[str] = Field(..., min_length=1, max_length=256)


class EmbedResult(BaseModel):
    vectors: list[list[float]]
    model: str
    dimension: int


class ChunkOut(BaseModel):
    index: int
    text: str
    char_start: int
    char_end: int
    char_count: int
    vector: list[float] | None = None


class ProcessResult(BaseModel):
    text: str
    page_count: int
    char_count: int
    detected_type: str
    chunks: list[ChunkOut]
    model: str
    dimension: int


# ----------------------------- extraction -----------------------------

def _extract_pdf(data: bytes) -> tuple[str, int]:
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        if doc.needs_pass:
            raise HTTPException(422, "PDF is password protected and cannot be read.")
        pages = [page.get_text() for page in doc]
    return "\n\n".join(pages), len(pages)


def _extract_docx(data: bytes) -> tuple[str, int]:
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Tables carry a lot of the content in clinical documents; paragraphs alone miss them.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts), 1


def _extract_plain(data: bytes) -> tuple[str, int]:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding), 1
        except UnicodeDecodeError:
            continue
    raise HTTPException(422, "Could not decode the file as text in any supported encoding.")


async def _read_and_extract(file: UploadFile) -> tuple[str, int, str]:
    data = await file.read()
    if not data:
        raise HTTPException(400, "File is empty.")
    if len(data) > MAX_BYTES:
        raise HTTPException(413, f"File exceeds the {MAX_BYTES // (1024 * 1024)}MB limit.")

    content_type = (file.content_type or "").split(";")[0].strip().lower()
    name = (file.filename or "").lower()

    if content_type in PDF_TYPES or name.endswith(".pdf"):
        text, pages = _extract_pdf(data)
        return text, pages, "pdf"
    if content_type in DOCX_TYPES or name.endswith(".docx"):
        text, pages = _extract_docx(data)
        return text, pages, "docx"
    if content_type in TEXT_TYPES or name.endswith((".txt", ".md", ".csv")):
        text, pages = _extract_plain(data)
        return text, pages, "text"

    raise HTTPException(
        415,
        f"Unsupported type '{content_type or 'unknown'}' for '{file.filename}'. "
        "Supported: PDF, DOCX, TXT.",
    )


# ----------------------------- endpoints -----------------------------

@app.get("/health", response_model=HealthResult)
def health() -> HealthResult:
    try:
        store = vectorstore.stats()
    except Exception as exc:
        store = {"available": False, "error": str(exc)}
    return HealthResult(
        status="healthy",
        pymupdf_version=pymupdf.__doc__.split()[1].rstrip(":"),
        embedding=embedding.model_info(),
        vector_store=store,
    )


@app.post("/extract-text", response_model=ExtractionResult)
async def extract_text(file: UploadFile = File(...)) -> ExtractionResult:
    text, pages, detected = await _read_and_extract(file)
    log.info("Extracted %s: %d chars from %d page(s)", file.filename, len(text), pages)
    return ExtractionResult(
        text=text, page_count=pages, char_count=len(text), detected_type=detected
    )


@app.post("/embed", response_model=EmbedResult)
def embed_texts(request: EmbedRequest) -> EmbedResult:
    """Embed arbitrary text. Used for search queries at retrieval time - the query
    must go through the same model as the documents, or the vectors are not
    comparable and every result is noise."""
    vectors = embedding.embed(request.texts)
    return EmbedResult(
        vectors=vectors, model=embedding.MODEL_NAME, dimension=embedding.EMBEDDING_DIM
    )


@app.post("/process", response_model=ProcessResult)
async def process(
    file: UploadFile = File(...),
    chunk_size: int = chunking.DEFAULT_CHUNK_SIZE,
    overlap: int = chunking.DEFAULT_OVERLAP,
    embed: bool = True,
) -> ProcessResult:
    """Extract, chunk and embed in one call - the ingestion path the .NET API uses.

    One request rather than three keeps the file bytes on this side of the wire;
    posting the document once beats shipping the extracted text back and forth."""
    text, pages, detected = await _read_and_extract(file)

    try:
        chunks = chunking.chunk_text(text, chunk_size=chunk_size, overlap=overlap)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc

    vectors: list[list[float] | None] = [None] * len(chunks)
    if embed and chunks:
        vectors = embedding.embed([c.text for c in chunks])

    log.info(
        "Processed %s: %d chars -> %d chunk(s)%s",
        file.filename, len(text), len(chunks), " (embedded)" if embed and chunks else "",
    )

    return ProcessResult(
        text=text,
        page_count=pages,
        char_count=len(text),
        detected_type=detected,
        chunks=[
            ChunkOut(
                index=c.index, text=c.text, char_start=c.char_start,
                char_end=c.char_end, char_count=c.char_count, vector=v,
            )
            for c, v in zip(chunks, vectors)
        ],
        model=embedding.MODEL_NAME,
        dimension=embedding.EMBEDDING_DIM,
    )


# ----------------------------- RAG endpoints -----------------------------

class IndexResult(BaseModel):
    document_id: str
    chunks_indexed: int
    char_count: int
    page_count: int
    detected_type: str
    # The caller persists this in Postgres. Returning it here means one round trip
    # rather than extracting the same file twice.
    text: str


class SearchRequest(BaseModel):
    query: str = Field(..., min_length=1, max_length=2000)
    limit: int = Field(5, ge=1, le=50)
    document_id: str | None = None
    score_threshold: float | None = None


class SearchHit(BaseModel):
    score: float
    document_id: str
    file_name: str
    document_type: str
    chunk_index: int
    text: str
    char_start: int
    char_end: int


class SearchResponse(BaseModel):
    query: str
    hits: list[SearchHit]
    model: str


@app.post("/index", response_model=IndexResult)
async def index_document(
    file: UploadFile = File(...),
    document_id: str = Form(...),
    document_type: str = Form("Report"),
    chunk_size: int = chunking.DEFAULT_CHUNK_SIZE,
    overlap: int = chunking.DEFAULT_OVERLAP,
) -> IndexResult:
    """Extract, chunk, embed and store - the full ingestion path.

    `document_id` comes from the caller (the Postgres row id), so the vector store
    and the relational record stay in step and a re-index overwrites cleanly."""
    text, pages, detected = await _read_and_extract(file)

    try:
        chunks = chunking.chunk_text(text, chunk_size=chunk_size, overlap=overlap)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc

    if not chunks:
        return IndexResult(document_id=document_id, chunks_indexed=0, char_count=len(text),
                           page_count=pages, detected_type=detected, text=text)

    vectors = embedding.embed([c.text for c in chunks])
    payload = [
        {"index": c.index, "text": c.text, "char_start": c.char_start,
         "char_end": c.char_end, "vector": v}
        for c, v in zip(chunks, vectors)
    ]

    try:
        stored = vectorstore.upsert_chunks(document_id, file.filename or "", document_type, payload)
    except Exception as exc:
        raise HTTPException(503, f"Vector store unavailable: {exc}") from exc

    return IndexResult(document_id=document_id, chunks_indexed=stored, char_count=len(text),
                       page_count=pages, detected_type=detected, text=text)


@app.post("/search", response_model=SearchResponse)
def search_chunks(request: SearchRequest) -> SearchResponse:
    """Top-K retrieval. The query is embedded here with the same model as the
    documents - accepting a caller-supplied vector would invite a mismatch that
    returns confident nonsense."""
    try:
        hits = vectorstore.search(
            request.query, limit=request.limit,
            document_id=request.document_id, score_threshold=request.score_threshold,
        )
    except Exception as exc:
        raise HTTPException(503, f"Vector store unavailable: {exc}") from exc

    return SearchResponse(
        query=request.query,
        hits=[SearchHit(**h) for h in hits],
        model=embedding.MODEL_NAME,
    )


@app.delete("/index/{document_id}")
def delete_index(document_id: str) -> dict:
    """Drop a document's chunks. Without this the index keeps citing text that no
    longer exists once a document is deleted."""
    try:
        vectorstore.delete_document(document_id)
    except Exception as exc:
        raise HTTPException(503, f"Vector store unavailable: {exc}") from exc
    return {"document_id": document_id, "deleted": True}
